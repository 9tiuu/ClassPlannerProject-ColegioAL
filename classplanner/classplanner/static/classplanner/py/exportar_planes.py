import unicodedata
from html import escape
from io import BytesIO
from pathlib import Path
import reportlab
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table as PDFTable,
    TableStyle as PDFTableStyle,
)
from openpyxl import load_workbook
from openpyxl import Workbook
from openpyxl.styles import Font, Border, Alignment, Side
from openpyxl.utils import get_column_letter
from django.db.models.query import QuerySet
from django.utils import timezone
from django.http import HttpResponse
from openpyxl.worksheet.table import Table
from core.middleware.constants import asigns_1_4_b, asigns_5_6_b, asigns_7_8_b, asigns_1_2_m, asigns_3_m, asigns_4_m

########## EXCEL 
titulo_font = Font(
    name='Arial',
    size=16,
    bold=True,
    italic=False,
    vertAlign=None,
    underline='none',
    strike=False,
)
subtitle_font = Font(
    name='Arial',
    size=11,
    bold=True,
    italic=False,
    vertAlign=None,
    underline='none',
    strike=False,
)
center_align = Alignment(
    horizontal='center',
    vertical='center',
    indent=0
)

plan_table_border = Border(
    left=Side(style='thin', color='000000'),
    right=Side(style='thin', color='000000'),
    top=Side(style='thin', color='000000'),
    bottom=Side(style='thin', color='000000'),
)

reportlab_fonts = Path(reportlab.__file__).parent / 'fonts'
pdfmetrics.registerFont(TTFont('Vera', str(reportlab_fonts / 'Vera.ttf')))
pdfmetrics.registerFont(TTFont('Vera-Bold', str(reportlab_fonts / 'VeraBd.ttf')))

MINEDUC_TEMPLATES = {
    '1°B': ('1°-4° Básico', asigns_1_4_b),
    '2°B': ('1°-4° Básico', asigns_1_4_b),
    '3°B': ('1°-4° Básico', asigns_1_4_b),
    '4°B': ('1°-4° Básico', asigns_1_4_b),
    '5°B': ('5°-6° Básico', asigns_5_6_b),
    '6°B': ('5°-6° Básico', asigns_5_6_b),
    '7°B': ('7°-8° Básico', asigns_7_8_b),
    '8°B': ('7°-8° Básico', asigns_7_8_b),
    '1°M': ('1°-2° Medio', asigns_1_2_m),
    '2°M': ('1°-2° Medio', asigns_1_2_m),
    '3°M': ('3° Medio', asigns_3_m),
    '4°M': ('4° Medio', asigns_4_m),
}


def _format_pdf_value(value):
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _exportar_planes_pdf(planes):
    planes_exportados = list(
        planes.select_related(
            'curso_id',
            'asignatura_id',
            'asignatura_id__plan_diferencial_id',
        )
    )
    cursos_por_grupo = {}
    horas_por_asignatura_y_curso = {}
    horas_plan_diferencial_por_curso = {}

    for plan in planes_exportados:
        curso_id = plan.curso_id_id
        nivel = plan.curso_id.nivel
        if nivel not in MINEDUC_TEMPLATES:
            raise ValueError(f'No hay una plantilla MINEDUC para el curso {nivel}')

        sheet_title, _ = MINEDUC_TEMPLATES[nivel]
        grupo = cursos_por_grupo.setdefault(sheet_title, {'courses': {}})
        grupo['courses'].setdefault(curso_id, str(plan.curso_id))

        key = (str(plan.asignatura_id.nombre).lower(), curso_id)
        horas_por_asignatura_y_curso[key] = (
            horas_por_asignatura_y_curso.get(key, 0) + plan.hrs_asignatura
        )

        if (
            nivel in ('3°M', '4°M')
            and plan.asignatura_id.plan_diferencial_id is not None
        ):
            horas_plan_diferencial_por_curso[curso_id] = (
                horas_plan_diferencial_por_curso.get(curso_id, 0)
                + plan.hrs_asignatura
            )

    grupos_ordenados = []
    grupos_vistos = set()
    for sheet_title, template in MINEDUC_TEMPLATES.values():
        if sheet_title in cursos_por_grupo and sheet_title not in grupos_vistos:
            grupos_ordenados.append((sheet_title, template))
            grupos_vistos.add(sheet_title)

    if not grupos_ordenados:
        grupos_ordenados = [('1°-4° Básico', asigns_1_4_b)]
        cursos_por_grupo['1°-4° Básico'] = {'courses': {}}

    page_width, _ = landscape(A4)
    pdf_buffer = BytesIO()
    document = SimpleDocTemplate(
        pdf_buffer,
        pagesize=landscape(A4),
        leftMargin=28,
        rightMargin=28,
        topMargin=28,
        bottomMargin=28,
        title='Planes de Estudio',
    )
    available_width = page_width - document.leftMargin - document.rightMargin
    subject_width, annual_width, weekly_width = 205, 60, 60
    course_width = 66
    course_columns_per_page = max(
        1,
        int(
            (available_width - subject_width - annual_width - weekly_width)
            // course_width
        ),
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'PlanTitle',
        parent=styles['Heading2'],
        fontName='Vera-Bold',
        fontSize=13,
        leading=16,
        textColor=colors.HexColor('#234FA8'),
        spaceAfter=8,
    )
    header_style = ParagraphStyle(
        'PlanHeader',
        parent=styles['BodyText'],
        fontName='Vera-Bold',
        fontSize=8,
        leading=10,
        alignment=1,
    )
    body_style = ParagraphStyle(
        'PlanCell',
        parent=styles['BodyText'],
        fontName='Vera',
        fontSize=8,
        leading=10,
    )
    centered_style = ParagraphStyle(
        'PlanCenteredCell',
        parent=body_style,
        alignment=1,
    )

    def paragraph(value, style):
        return Paragraph(escape(_format_pdf_value(value)).replace('\n', '<br/>'), style)

    story = []
    for sheet_title, template in grupos_ordenados:
        cursos = list(cursos_por_grupo[sheet_title]['courses'].items())
        course_chunks = [
            cursos[index:index + course_columns_per_page]
            for index in range(0, len(cursos), course_columns_per_page)
        ] or [[]]

        for course_chunk in course_chunks:
            if story:
                story.append(PageBreak())
            story.append(paragraph(f'Plan de Estudios - {sheet_title}', title_style))
            story.append(Spacer(1, 4))

            course_labels = [
                course_name[:-4].replace('B', ' básico').replace('M', ' medio')
                for _, course_name in course_chunk
            ]
            header_row = [
                paragraph('Plan de Estudios', header_style),
                paragraph('Horas pedagógicas\n(MINEDUC)', header_style),
                '',
                paragraph('Horas pedagógicas semanales', header_style),
                *([''] * (len(course_labels) - 1)),
            ]
            second_header_row = [
                '',
                paragraph('Anuales', header_style),
                paragraph('Semanales', header_style),
                *[paragraph(label, header_style) for label in course_labels],
            ]
            if not course_chunk:
                second_header_row.append('')

            table_data = [header_row, second_header_row]
            total_anual = 0
            total_semanal = 0
            for asign_nombre, horas_semanales_mineduc in template.items():
                hrs_anuales = (
                    horas_semanales_mineduc * 38
                    if isinstance(horas_semanales_mineduc, (int, float))
                    else horas_semanales_mineduc
                )
                if isinstance(hrs_anuales, (int, float)):
                    total_anual += hrs_anuales
                if isinstance(horas_semanales_mineduc, (int, float)):
                    total_semanal += horas_semanales_mineduc

                row = [
                    paragraph(asign_nombre, body_style),
                    paragraph(hrs_anuales, centered_style),
                    paragraph(horas_semanales_mineduc, centered_style),
                ]
                for curso_id, _ in course_chunk:
                    horas = (
                        horas_plan_diferencial_por_curso.get(curso_id, 0)
                        if asign_nombre == 'Plan Diferenciado'
                        else horas_por_asignatura_y_curso.get(
                            (asign_nombre.lower(), curso_id),
                            0,
                        )
                    )
                    row.append(paragraph(horas, centered_style))
                if not course_chunk:
                    row.append('')
                table_data.append(row)

            total_row = [
                paragraph('Total tiempo escolar', header_style),
                paragraph(total_anual, header_style),
                paragraph(total_semanal, header_style),
            ]
            for curso_id, _ in course_chunk:
                total = sum(
                    horas_por_asignatura_y_curso.get(
                        (asign_nombre.lower(), curso_id),
                        0,
                    )
                    for asign_nombre in template
                )
                total += horas_plan_diferencial_por_curso.get(curso_id, 0)
                total_row.append(paragraph(total, header_style))
            if not course_chunk:
                total_row.append('')
            table_data.append(total_row)

            column_widths = [
                subject_width,
                annual_width,
                weekly_width,
                *([course_width] * (len(course_chunk) or 1)),
            ]
            table = PDFTable(
                table_data,
                colWidths=column_widths,
                repeatRows=2,
                hAlign='LEFT',
            )
            table_style = [
                ('GRID', (0, 0), (-1, -1), 0.5, colors.black),
                ('SPAN', (0, 0), (0, 1)),
                ('SPAN', (1, 0), (2, 0)),
                ('BACKGROUND', (0, 0), (-1, 1), colors.HexColor('#E8EEF8')),
                ('ALIGN', (0, 0), (-1, 1), 'CENTER'),
                ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                ('ALIGN', (0, 2), (0, -1), 'LEFT'),
                ('ALIGN', (1, 2), (-1, -1), 'CENTER'),
                ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor('#E8EEF8')),
                ('FONTNAME', (0, -1), (-1, -1), 'Vera-Bold'),
                ('LEFTPADDING', (0, 0), (-1, -1), 4),
                ('RIGHTPADDING', (0, 0), (-1, -1), 4),
                ('TOPPADDING', (0, 0), (-1, -1), 4),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
            ]
            table_style.append(('SPAN', (3, 0), (-1, 0)))
            table.setStyle(PDFTableStyle(table_style))
            story.append(table)

    document.build(story)
    response = HttpResponse(pdf_buffer.getvalue(), content_type='application/pdf')
    response['Content-Disposition'] = 'attachment; filename=PlanesDeEstudio.pdf'
    return response


def exportar_planes(planes=QuerySet, export_to=str):
    if export_to == 'excel':
        planes_exportados = list(planes.select_related('curso_id', 'asignatura_id'))
        cursos_por_grupo = {}
        horas_por_asignatura_y_curso = {}
        horas_plan_diferencial_por_curso = {}

        for plan in planes_exportados:
            curso_id = plan.curso_id_id
            nivel = plan.curso_id.nivel
            if nivel not in MINEDUC_TEMPLATES:
                raise ValueError(f'No hay una plantilla MINEDUC para el curso {nivel}')

            sheet_title, _ = MINEDUC_TEMPLATES[nivel]
            grupo = cursos_por_grupo.setdefault(
                sheet_title,
                {'courses': {}},
            )
            grupo['courses'].setdefault(curso_id, str(plan.curso_id))

            key = (str(plan.asignatura_id.nombre).lower(), curso_id)
            horas_por_asignatura_y_curso[key] = (
                horas_por_asignatura_y_curso.get(key, 0) + plan.hrs_asignatura
            )

            if (nivel == '4°M' or nivel == '3°M') and plan.asignatura_id.plan_diferencial_id != None:
                horas_plan_diferencial_por_curso[curso_id] = (
                    horas_plan_diferencial_por_curso.get(curso_id, 0)
                    + plan.hrs_asignatura
                )
                print("horas_plan_diferencial_por_curso: ",horas_plan_diferencial_por_curso)
                print("plan.asignatura_id.plan_diferencial_id.nombre: ",plan.asignatura_id.plan_diferencial_id.nombre)


        wb = Workbook(write_only=False)
        grupos_ordenados = []
        grupos_vistos = set()
        for sheet_title, template in MINEDUC_TEMPLATES.values():
            if sheet_title in cursos_por_grupo and sheet_title not in grupos_vistos:
                grupos_ordenados.append((sheet_title, template))
                grupos_vistos.add(sheet_title)

        if not grupos_ordenados:
            grupos_ordenados = [('1°-4° Básico', asigns_1_4_b)]
            cursos_por_grupo['1°-4° Básico'] = {
                'courses': {},
            }

        for index, (sheet_title, template) in enumerate(grupos_ordenados):
            group = cursos_por_grupo[sheet_title]
            cursos_ids = list(group['courses'])
            worksheet = wb.active if index == 0 else wb.create_sheet()
            total_semanal = 0
            total_anual = 0
            create_file(
                wb,
                [group['courses'][cur_id] for cur_id in cursos_ids],
                worksheet=worksheet,
                sheet_title=sheet_title,
            )

            for row, (asign_nombre, horas_semanales_mineduc) in enumerate(
                template.items(),
                start=3,
            ):
                hrs_anuales = (
                    horas_semanales_mineduc * 38
                    if isinstance(horas_semanales_mineduc, (int, float))
                    else horas_semanales_mineduc
                )
                worksheet.cell(row=row, column=1, value=asign_nombre)
                worksheet.cell(row=row, column=2, value=hrs_anuales)
                total_anual += hrs_anuales
                worksheet.cell(row=row, column=3, value=horas_semanales_mineduc)
                total_semanal += horas_semanales_mineduc
                for column, cur_id in enumerate(cursos_ids, start=4):
                    if asign_nombre == 'Plan Diferenciado':
                        worksheet.cell(
                            row=row,
                            column=column,
                            value=horas_plan_diferencial_por_curso.get(cur_id, 0),
                        )
                    else:
                        worksheet.cell(
                            row=row,
                            column=column,
                            value=horas_por_asignatura_y_curso.get(
                                (asign_nombre.lower(), cur_id),
                                0,
                            ),
                        )

            total_row = len(template) + 3
            worksheet.cell(total_row, column=1, value='Total tiempo escolar')
            worksheet.cell(total_row, column=2, value=total_anual)
            worksheet.cell(total_row, column=3, value=total_semanal)

            for column, cur_id in enumerate(cursos_ids, start=4):
                total = sum(
                    horas_por_asignatura_y_curso.get(
                        ((asign_nombre).lower(), cur_id),
                        0
                    )
                    for asign_nombre in template
                )
                total += horas_plan_diferencial_por_curso.get(cur_id, 0)
                print("total als: ",total)
                worksheet.cell(total_row, column=column, value=total)
                worksheet.cell(total_row, column=column).number_format = '#0'

            last_column = 3 + len(cursos_ids)
            for row in worksheet.iter_rows(
                min_row=1,
                max_row=total_row,
                min_col=1,
                max_col=max(3, last_column),
            ):
                for cell in row:
                    cell.border = plan_table_border
                    cell.font = Font(
                        name='Arial',
                        size=11,
                        bold=(
                            cell.row <= 2
                            or cell.row == total_row
                            or (cell.column >= 4 and cell.row >= 3)
                        ),
                        color='000000',
                    )
                    if cell.column >= 4 and cell.row >= 3:
                        cell.number_format = '#0'
                    cell.alignment = Alignment(
                        horizontal='left'
                        if cell.column == 1 and cell.row >= 3
                        else 'center',
                        vertical='center',
                    )
                    if cell.row >= 3:
                        worksheet.row_dimensions[cell.row].height = 17

        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename=PlanesDeEstudio.xlsx'
        wb.save(response)
        return response
    
    elif export_to == 'pdf':
        return _exportar_planes_pdf(planes)

    def create_file(
        wb=None,
        curso_names=None,
        worksheet=None,
        sheet_title='PlanesDeEstudio',
    ):
        if wb is None:
            wb = Workbook()
        if curso_names is None:
            curso_names = ('curso 1', 'curso 2')

        ws = worksheet or wb.active
        ws.title = sheet_title

        curso_names = list(curso_names)
        last_curso_column = max(4, 3 + len(curso_names))
        ws.merge_cells('A1:A2')
        ws.merge_cells('B1:C1')
        ws.merge_cells(
            start_row=1,
            start_column=4,
            end_row=1,
            end_column=last_curso_column,
        )
        ws['A1'] = 'Plan de Estudios'
        ws['B1'] = 'Horas pedagógicas\n(MINEDUC)'
        ws['B2'] = 'Anuales'
        ws['C2'] = 'Semanales'
        ws.cell(row=1, column=4, value='Horas pedagógicas semanales')

        for column, curso_name in enumerate(curso_names, start=4):
            ws.cell(row=2, column=column, value=curso_name[:-4].replace("B"," básico").replace("M"," medio"))

        ws.column_dimensions['A'].width = 38
        ws.column_dimensions['B'].width = 13
        ws.column_dimensions['C'].width = 13
        for column in range(4, last_curso_column + 1):
            ws.column_dimensions[get_column_letter(column)].width = 18
        for row in ws.iter_rows(
            min_row=1,
            max_row=2,
            min_col=1,
            max_col=last_curso_column,
        ):
            for cell in row:
                cell.border = plan_table_border
                cell.font = subtitle_font
                cell.alignment = Alignment(
                    horizontal='center',
                    vertical='center',
                    wrap_text=True,
                )


        ws['B1'].alignment = Alignment(
            horizontal='center',
            vertical='center',
            wrap_text=True,
        )
        ws.row_dimensions[1].height = 36
        ws.row_dimensions[2].height = 18
